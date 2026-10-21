# Read the All Theory workbook and turn it into the JSON the app embeds. Prints any problems it finds.
import sys,json,re
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],data_only=True);probs=[]
GR=['10th_Kup','9th_Kup','8th_Kup','7th_Kup','6th_Kup','5th_Kup','4th_Kup','3rd_Kup','2nd_Kup','1st_Kup','1st_Dan','2nd_Dan','3rd_Dan','4th_Dan','5th_Dan']
def rows(t):
    ws=wb[t];h=[str(c.value or '').strip() for c in ws[1]]
    need={'Terminology':['ID','Grade','Belt','English','Korean','Pronounced','MP3'],'Patterns':['ID','Grade','Belt','Pattern','Moves','Meaning','Pronounced','MP3'],'Belt meanings':['ID','Belt','First grade','Meaning','Pronounced','MP3']}.get(t,[])
    for n in need:
        if n not in h:probs.append(f'{t}: header "{n}" is missing (headers are {h})')
    for r in ws.iter_rows(min_row=2,values_only=True):
        if not any(v not in (None,'') for v in r):continue
        yield {h[i]:(str(v).strip() if v is not None else '') for i,v in enumerate(r) if i<len(h)}
def col(d,name):
    for k in d:
        if k.lower().startswith(name.lower()):return d[k]
    return ''
def secs(v,where):
    v=str(v).strip()
    if not v:return None
    if re.fullmatch(r'\d+(\.\d+)?',v):
        f=float(v)
        if 0<f<1:return round(f*86400)          # Excel turned 0:21 into a time of day
        return int(f)
    m=re.fullmatch(r'(?:(\d+):)?(\d{1,2}):(\d{2})(?::00)?',v)
    if m:
        h=int(m.group(1) or 0);a=int(m.group(2));b=int(m.group(3))
        return h*3600+a*60+b
    probs.append(f'{where}: time "{v}" not understood');return None
def vid(u,where):
    if not u:return None
    m=re.search(r'(?:[?&]v=|youtu\.be/|/shorts/|/embed/|/live/)([A-Za-z0-9_-]{11})',u)
    if not m:probs.append(f'{where}: link "{u}" has no single video in it');return None
    return m.group(1)
out={'terms':[],'patterns':[],'belts':[]};seen=set()
def chk_id(i,where):
    if not i:probs.append(f'{where}: missing ID')
    elif i in seen:probs.append(f'{where}: ID {i} used twice')
    seen.add(i)
for d in rows('Terminology'):
    w=f"Terminology {d.get('ID')}";chk_id(d.get('ID'),w)
    if d['Grade'] not in GR:probs.append(f'{w}: unknown grade "{d["Grade"]}"')
    if not d['English'] or not d['Korean']:probs.append(f'{w}: English or Korean is empty')
    out['terms'].append({'id':d['ID'],'g':d['Grade'],'en':d['English'],'kr':d['Korean'],'say':d.get('Pronounced',''),'mp3':d.get('MP3','')})
for d in rows('Patterns'):
    w=f"Patterns {d.get('ID')} ({d.get('Pattern')})";chk_id(d.get('ID'),w)
    if d['Grade'] not in GR:probs.append(f'{w}: unknown grade "{d["Grade"]}"')
    v=vid(col(d,'Video link'),w);s=secs(col(d,'Video start'),w);e=secs(col(d,'Video end'),w)
    if v and s is not None and e is not None and e<=s:probs.append(f'{w}: video end {col(d,"Video end")} is not after start {col(d,"Video start")}')
    if not v and (s or e):probs.append(f'{w}: has times but no link')
    out['patterns'].append({'id':d['ID'],'g':d['Grade'],'name':d['Pattern'],'moves':int(float(d['Moves'])) if d['Moves'] else None,'meaning':d['Meaning'],
        'say':d.get('Pronounced',''),'mp3':d.get('MP3',''),'video':{'id':v,'start':s or 0,'end':e} if v else None})
for d in rows('Belt meanings'):
    w=f"Belt meanings {d.get('ID')}";chk_id(d.get('ID'),w)
    out['belts'].append({'id':d['ID'],'belt':d['Belt'],'g':d['First grade'],'meaning':d['Meaning'],'say':d.get('Pronounced',''),'mp3':d.get('MP3','')})
# Questions tab (optional): Grade is All or a grade; answers and wrong answers separated by |
out['questions']=[]
if 'Questions' in wb.sheetnames:
    split=lambda v:[x.strip() for x in str(v or '').split('|') if x.strip()]
    for d in rows('Questions'):
        w=f"Questions {d.get('ID')}";chk_id(d.get('ID'),w)
        g=d.get('Grade','')
        if g!='All' and g not in GR:probs.append(f'{w}: unknown grade "{g}"')
        i=d.get('ID','');exp='Q_All_' if g=='All' else 'Q_'+g.replace('_','-')+'_'
        if i and not i.startswith(exp):probs.append(f'{w}: ID does not match its grade (expected {exp}..)')
        if not d.get('Question') or not split(d.get('Answer')):probs.append(f'{w}: Question or Answer is empty');continue
        ans=split(d.get('Answer'));wr=[x for x in split(d.get('Wrong answers')) if x.lower() not in [a.lower() for a in ans]]
        out['questions'].append({'id':i,'g':g,'q':d['Question'],'a':ans,'w':wr,
            'qsay':col(d,'Question pronounced') or d.get('Pronounced',''),'qmp3':col(d,'Question MP3') or d.get('MP3',''),
            'asay':col(d,'Answer pronounced'),'amp3':col(d,'Answer MP3')})
json.dump(out,open(sys.argv[2],'w'),ensure_ascii=False)
print(len(out.get('questions',[])),'questions,',len(out['terms']),'terms,',len(out['patterns']),'patterns,',len(out['belts']),'belts,',sum(1 for p in out['patterns'] if p['video']),'videos')
print('PROBLEMS:' if probs else 'No problems found.');[print(' -',p) for p in probs]
print('No video:',[p['name'] for p in out['patterns'] if not p['video']])
